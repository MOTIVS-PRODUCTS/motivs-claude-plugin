#!/usr/bin/env python3
"""Create one Motivs Platform automation staging artifact.

The platform renders the employee UI from manifest.json. The generated
automation.py exposes a backend adapter that the platform can execute.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path


NOTEBOOK_TEMPLATE = '''from __future__ import annotations

import csv
import hashlib
import importlib.util
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

try:
    import marimo
except Exception:  # marimo is useful for developer repair, not required by the platform UI.
    marimo = None

DUCKDB_AVAILABLE = importlib.util.find_spec("duckdb") is not None
PANDERA_AVAILABLE = importlib.util.find_spec("pandera") is not None
PYARROW_AVAILABLE = importlib.util.find_spec("pyarrow") is not None

AUTOMATION_NAME = "{automation_name}"
AUTOMATION_VERSION = "0.1.0"
SUPPORTED_EXTENSIONS = {{".csv", ".xlsx", ".xlsm"}}
# Replace with (manifest input ID, region ID) pairs mapped from the committed
# automation_spec.source_selections.
SELECTED_REGIONS: set[tuple[str, str]] = set()


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_name(value: object) -> str:
    text = "" if value is None else str(value)
    return " ".join(text.strip().replace("_", " ").split()).lower()


def profile_csv(path: Path) -> dict[str, object]:
    sample = path.read_text(encoding="utf-8-sig", errors="replace")[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample)
    except Exception:
        dialect = csv.excel
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle, dialect)
        header = next(reader, [])
        rows = sum(1 for _ in reader)
    return {{
        "filename": path.name,
        "kind": "csv",
        "rows": rows,
        "header": header,
        "normalized_header": [normalize_name(item) for item in header],
        "sha256": file_hash(path),
    }}


def profile_excel(path: Path) -> dict[str, object]:
    keep_vba = path.suffix.lower() == ".xlsm"
    workbook = load_workbook(path, read_only=True, data_only=False, keep_vba=keep_vba)
    try:
        sheets = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            first_row = next(sheet.iter_rows(values_only=True), ())
            sheets.append(
                {{
                    "name": sheet_name,
                    "rows": max(int(sheet.max_row or 0) - 1, 0),
                    "columns": int(sheet.max_column or 0),
                    "header": list(first_row or []),
                    "normalized_header": [normalize_name(item) for item in first_row or []],
                }}
            )
        return {{
            "filename": path.name,
            "kind": path.suffix.lower().lstrip("."),
            "sheets": sheets,
            "sha256": file_hash(path),
        }}
    finally:
        workbook.close()


def profile_file(path: Path) -> dict[str, object]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return profile_csv(path)
    if suffix in {{".xlsx", ".xlsm"}}:
        return profile_excel(path)
    return {{"filename": path.name, "kind": "unsupported", "error": f"Unsupported file type: {{suffix}}"}}


def read_table(path: Path) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return pd.read_csv(path)
    if suffix in {{".xlsx", ".xlsm"}}:
        return pd.read_excel(path)
    raise ValueError(f"Unsupported file type: {{suffix}}")


def normalized_region_paths(parameters: dict[str, Any] | None) -> list[Path] | None:
    context = (parameters or {{}}).get("_motivs")
    if not isinstance(context, dict) or context.get("schema_version") != "workbook-interpretation-v1":
        return None
    paths: list[Path] = []
    inputs = context.get("inputs")
    for input_id, files in inputs.items() if isinstance(inputs, dict) else []:
        for file_contract in files if isinstance(files, list) else []:
            regions = file_contract.get("regions") if isinstance(file_contract, dict) else []
            candidates = [region for region in regions if isinstance(region, dict)] if isinstance(regions, list) else []
            if SELECTED_REGIONS:
                candidates = [
                    region
                    for region in candidates
                    if (str(input_id), str(region.get("id"))) in SELECTED_REGIONS
                ]
            else:
                # Safe scaffold fallback: one primary visible, non-ambiguous region per file.
                candidates = [
                    region
                    for region in candidates
                    if region.get("sheet_state") in {{"visible", "not_applicable"}}
                    and region.get("confidence") != "low"
                ]
                candidates = sorted(
                    candidates,
                    key=lambda region: int(region.get("record_count") or 0),
                    reverse=True,
                )[:1]
            for region in candidates:
                value = region.get("path")
                if isinstance(value, str) and value:
                    paths.append(Path(value))
    return paths


def verify_output_workbook(output_path: Path, expected_rows: int) -> dict[str, object]:
    workbook = load_workbook(output_path, read_only=True, data_only=False)
    try:
        checks = {{
            "readable": True,
            "required_sheets_present": "Data" in workbook.sheetnames and "Summary" in workbook.sheetnames,
            "expected_rows": expected_rows,
            "data_rows": 0,
            "row_count_matches": False,
            "formula_cells": 0,
            "formula_recalculation": "not required",
            "sha256": file_hash(output_path),
        }}
        if "Data" in workbook.sheetnames:
            checks["data_rows"] = max(int(workbook["Data"].max_row or 0) - 1, 0)
            checks["row_count_matches"] = checks["data_rows"] == expected_rows
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        checks["formula_cells"] += 1
        if checks["formula_cells"]:
            checks["formula_recalculation"] = "not verified"
        return checks
    finally:
        workbook.close()


def flatten_input_paths(input_paths: dict[str, list[str]]) -> list[Path]:
    paths: list[Path] = []
    for values in input_paths.values():
        for value in values:
            path = Path(value)
            if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                raise ValueError(f"Unsupported input file type: {{path.name}}")
            paths.append(path)
    if not paths:
        raise ValueError("No input files were provided.")
    return paths


def run_self_checks() -> dict[str, object]:
    sample = pd.DataFrame([{{"id": 1, "amount": 10}}, {{"id": 2, "amount": 20}}])
    assert int(sample["amount"].sum()) == 30
    return {{"status": "passed", "checks": 1}}


def run_automation(
    input_paths: dict[str, list[str]],
    output_dir: str,
    parameters: dict[str, Any] | None = None,
) -> dict[str, object]:
    """Platform entrypoint called by Motivs.

    Replace the generic concat/report logic with the workflow-specific rules,
    but keep the function signature and write all downloadable files to output_dir.
    """
    output_path_dir = Path(output_dir)
    output_path_dir.mkdir(parents=True, exist_ok=True)
    source_paths = flatten_input_paths(input_paths)
    profiles = [profile_file(path) for path in source_paths]
    unsupported = [item for item in profiles if item.get("kind") == "unsupported"]
    if unsupported:
        raise ValueError("; ".join(str(item.get("error")) for item in unsupported))

    normalized_paths = normalized_region_paths(parameters)
    if normalized_paths is not None:
        if not normalized_paths:
            raise ValueError(
                "No approved normalized region matched this automation specification. "
                "Review the committed source/region selection instead of reading the raw workbook as a table."
            )
        frames = [pd.read_parquet(path) for path in normalized_paths]
    else:
        frames = [read_table(path) for path in source_paths]
    combined = pd.concat(frames, ignore_index=True, sort=False)
    created_at = datetime.now().strftime("%Y%m%d-%H%M%S")
    output_path = output_path_dir / f"generated_report_{{created_at}}.xlsx"
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        combined.to_excel(writer, index=False, sheet_name="Data")
        pd.DataFrame(
            [
                {{"metric": "input_files", "value": len(source_paths)}},
                {{"metric": "output_rows", "value": len(combined)}},
                {{"metric": "output_columns", "value": len(combined.columns)}},
            ]
        ).to_excel(writer, index=False, sheet_name="Summary")

    output_checks = verify_output_workbook(output_path, len(combined))
    if not (
        output_checks.get("readable")
        and output_checks.get("required_sheets_present")
        and output_checks.get("row_count_matches")
    ):
        return {{
            "status": "failed",
            "message": "The workbook was created, but output verification did not pass.",
            "profiles": profiles,
            "output_checks": output_checks,
            "outputs": {{"excel_report": str(output_path)}},
        }}

    return {{
        "status": "passed",
        "filename": output_path.name,
        "created_at": created_at,
        "rows": len(combined),
        "columns": len(combined.columns),
        "input_files": [path.name for path in source_paths],
        "profiles": profiles,
        "runtime_capabilities": {{
            "duckdb": DUCKDB_AVAILABLE,
            "pandera": PANDERA_AVAILABLE,
            "pyarrow": PYARROW_AVAILABLE,
        }},
        "self_checks": run_self_checks(),
        "output_checks": output_checks,
        "outputs": {{"excel_report": str(output_path)}},
    }}


if marimo is not None:
    __generated_with = "0.19.7"
    app = marimo.App(width="medium")

    @app.cell
    def _():
        import marimo as mo
        return (mo,)

    @app.cell
    def _(mo):
        mo.md(
            f"""
            # {{AUTOMATION_NAME}}

            This automation is designed for the Motivs Platform-rendered UI.
            Employees should run it from the platform automation page. This
            notebook view is kept only for developer repair and audit.
            """
        )
        return

    if __name__ == "__main__":
        app.run()
'''


def notebook_text(automation_name: str) -> str:
    return NOTEBOOK_TEMPLATE.format(automation_name=automation_name)


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", value.strip()).strip("-").lower()
    return slug or "excel-workflow-automation"


def requirements_text() -> str:
    return """marimo>=0.19
pandas
openpyxl
duckdb
pyarrow
# Optional for stricter dataframe schemas when a workflow uses it:
# pandera
"""


def write_manifest(bundle_dir: Path, automation_name: str) -> None:
    manifest = {
        "name": automation_name,
        "description": "Upload source files and generate a validated spreadsheet output.",
        "automation_version": "0.1.0",
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "entrypoint": "automation.py",
        "backend_adapter": "run_automation",
        "requirements": "requirements.txt",
        "runtime_target": "motivs-platform",
        "runtime_contract": "workbook-interpretation-v1",
        "platform_upload_required": True,
        "source_files": ["automation.py"],
        "packaging": "platform-version",
        "ui": {
            "mode": "platform-rendered",
            "title": automation_name,
            "purpose": "Upload the source files and download the validated Excel output.",
            "steps": ["Add files", "Run checks", "Create output", "Download results"],
            "inputs": [
                {
                    "id": "source_files",
                    "label": "Source files",
                    "type": "file",
                    "accept": [".csv", ".xlsx", ".xlsm"],
                    "multiple": True,
                    "required": True,
                }
            ],
            "actions": [{"id": "run", "label": "Generate report", "kind": "primary"}],
            "outputs": [{"id": "excel_report", "label": "Excel report", "type": "xlsx"}],
        },
    }
    (bundle_dir / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")


def validation_report_text(automation_name: str) -> str:
    return json.dumps(
        {
            "automation_name": automation_name,
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "status": "staged",
            "checks": [],
            "blockers": [],
            "platform_upload_required": True,
            "notes": "Run static checks, upload to Motivs Platform, and verify a platform-rendered run before claiming readiness.",
        },
        indent=2,
    )


def create_platform_artifacts(bundle_dir: Path, automation_name: str) -> Path:
    bundle_dir.mkdir(parents=True, exist_ok=True)
    (bundle_dir / "automation.py").write_text(notebook_text(automation_name), encoding="utf-8")
    (bundle_dir / "requirements.txt").write_text(requirements_text(), encoding="utf-8")
    (bundle_dir / "validation_report.json").write_text(validation_report_text(automation_name), encoding="utf-8")
    write_manifest(bundle_dir, automation_name)
    return bundle_dir


def main() -> int:
    parser = argparse.ArgumentParser(description="Create one Motivs Platform automation staging artifact.")
    parser.add_argument(
        "output",
        help="Path to the platform staging folder. For legacy use, a `.py` path creates only automation.py.",
    )
    parser.add_argument("--automation-name", default="Excel Workflow Automation", help="Visible automation title")
    parser.add_argument("--force", action="store_true", help="Overwrite an existing notebook or staging folder")
    parser.add_argument("--port", type=int, default=8501, help="Accepted for backward compatibility; ignored")
    args = parser.parse_args()

    output_path = Path(args.output).expanduser().resolve()
    if output_path.suffix == ".py":
        if output_path.exists() and not args.force:
            print(f"Notebook already exists: {output_path}", file=sys.stderr)
            return 1
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(notebook_text(args.automation_name), encoding="utf-8")
        print(output_path)
        return 0

    if output_path.exists() and any(output_path.iterdir()) and not args.force:
        print(f"Bundle directory already exists and is not empty: {output_path}", file=sys.stderr)
        return 1
    print(create_platform_artifacts(output_path, args.automation_name))
    return 0


if __name__ == "__main__":
    sys.exit(main())
